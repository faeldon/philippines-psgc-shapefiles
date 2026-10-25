#   ---------------------------------------------------------------------------------
#   Copyright (c) James Faeldon. All rights reserved.
#   Licensed under the MIT License. See LICENSE in project root for information.
#   ---------------------------------------------------------------------------------
"""PSA's PSGC publications, one folder each under data/psgc/<id>/ (CONTRACT 1).

PSA's downloads stay exactly as published in data/psgc/raw/<id>/, which git ignores. What is
committed is each publication's manifest.toml (every raw file's PSA URL and SHA-256, and the
datafile's column layout) and the extract of the raw files: psgc.csv, national-summary.csv,
prov-summary.csv and summary-of-changes.csv. extract() writes them from the raw files and check()
proves they still match; the pipeline and the diff read only the committed files.
"""
from __future__ import annotations

import csv
import filecmp
import hashlib
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .release import Table, closed, read_toml

# Publications live here, relative to the working directory the pipeline runs from.
DATA_PSGC = Path("data/psgc")

# The PSGC datafile columns the levels read; a layout must name each and keep it.
PSGC_REQUIRED_COLUMNS = ("psgc_code", "name", "corr_code", "geo_level")

# The sheets each raw file role is extracted from, and the file each one becomes.
SHEETS = {
    "datafile": {"psgc": "psgc.csv", "national_summary": "national-summary.csv", "prov_summary": "prov-summary.csv"},
    "summary_of_changes": {"log": "summary-of-changes.csv"},
}

# The Summary of Changes' columns, as its header row names them.
SOC_COLUMNS = (
    "Region/Province/Municipal/Bgy. Name",
    "Unit Type",
    "New Code",
    "Mother Unit/Old Name",
    "Old Code",
    "Description/Legal basis",
    "Remarks",
)


class PsgcDataError(ValueError):
    """A publication whose manifest, raw files or extract can't be used."""


@dataclass(frozen=True)
class AddedRow:
    """A PSGC row the datafile lacks, added to the table."""

    code: int
    name: str
    geo_level: str
    note: str = ""


@dataclass(frozen=True)
class Layout:
    """The datafile's column names by position, the codes its NCR district rows lack, and rows it lacks."""

    columns: tuple[str, ...]
    drop: tuple[str, ...] = ()
    district_codes: tuple[tuple[str, int], ...] = ()
    add_rows: tuple[AddedRow, ...] = ()


@dataclass(frozen=True)
class RawFile:
    """One file PSA published with the announcement, and the sheets extracted from it."""

    role: str
    name: str
    url: str
    sha256: str
    sheets: dict[str, str]
    sections: tuple[str, ...] = ()  # a Summary of Changes: the log sections that are this publication's


@dataclass(frozen=True)
class Manifest:
    """One PSA publication: its raw files and its datafile layout."""

    id: str
    folder: Path
    as_of: str
    announcement: str
    files: tuple[RawFile, ...]
    layout: Layout
    note: str = ""

    def file(self, role: str) -> RawFile:
        """The raw file with this role."""
        return next(f for f in self.files if f.role == role)

    @property
    def soc_sections(self) -> tuple[str, ...]:
        """The Summary of Changes sections that are this publication's."""
        return self.file("summary_of_changes").sections


def _layout(t: Table) -> Layout:
    columns = tuple(t.take("columns", "strs"))
    drop = tuple(t.take("drop", "strs", []))
    repeated = sorted({c for c in columns if columns.count(c) > 1})
    if repeated:
        t.fail(t.key("columns"), f"{repeated[0]!r} named more than once")
    for name in PSGC_REQUIRED_COLUMNS:
        if name not in columns:
            t.fail(t.key("columns"), f"must name {name!r}")
    for name in drop:
        if name not in columns:
            t.fail(t.key("drop"), f"{name!r} is not in columns")
        if name in PSGC_REQUIRED_COLUMNS:
            t.fail(t.key("drop"), f"{name!r} is read by the levels and can't be dropped")
    add_rows = []
    for i, r in enumerate(t.tables("add_rows", [])):
        row = closed(
            r,
            AddedRow(
                code=r.take("code", "int"),
                name=r.take("name", "str"),
                geo_level=r.take("geo_level", "str"),
                note=r.take("note", "str", ""),
            ),
        )
        if row.code in [a.code for a in add_rows]:
            t.fail(f"{t.key('add_rows')}[{i}].code", f"duplicate code {row.code}")
        add_rows.append(row)
    district_codes = tuple(
        closed(d, (d.take("name", "str"), d.take("code", "int"))) for d in t.tables("district_codes", [])
    )
    return closed(t, Layout(columns=columns, drop=drop, district_codes=district_codes, add_rows=tuple(add_rows)))


def _raw_file(t: Table) -> RawFile:
    role = t.take("role", "str")
    if role not in SHEETS:
        t.fail(t.key("role"), f"must be one of {', '.join(SHEETS)}, not {role!r}")
    sha256 = t.take("sha256", "str")
    if not re.fullmatch(r"[0-9a-f]{64}", sha256):
        t.fail(t.key("sha256"), "must be 64 lowercase hex digits")
    s = t.table("sheets")
    sheets = closed(s, {key: s.take(key, "str") for key in SHEETS[role]})
    sections = tuple(t.take("sections", "strs")) if role == "summary_of_changes" else ()
    return closed(
        t,
        RawFile(
            role=role,
            name=t.take("name", "str"),
            url=t.take("url", "str"),
            sha256=sha256,
            sheets=sheets,
            sections=sections,
        ),
    )


def load_manifest(pub_id: str, root=None) -> Manifest:
    """Read and check data/psgc/<pub_id>/manifest.toml (under root, DATA_PSGC by default)."""
    folder = Path(DATA_PSGC if root is None else root) / pub_id
    top = read_toml(folder / "manifest.toml", PsgcDataError)
    manifest = Manifest(
        id=top.take("id", "str"),
        folder=folder,
        as_of=top.take("as_of", "str"),
        announcement=top.take("announcement", "str"),
        files=tuple(_raw_file(f) for f in top.tables("files")),
        layout=_layout(top.table("layout")),
        note=top.take("note", "str", ""),
    )
    top.close()
    if manifest.id != pub_id:
        top.fail("id", f"{manifest.id!r} doesn't match its folder {pub_id!r}")
    roles = sorted(f.role for f in manifest.files)
    if roles != sorted(SHEETS):
        top.fail("files", f"must hold one file of each role ({', '.join(sorted(SHEETS))}), not {roles}")
    return manifest


def publications(root=None) -> list[str]:
    """The ids of every publication folder with a manifest, in order."""
    root = Path(DATA_PSGC if root is None else root)
    return sorted(p.parent.name for p in root.glob("*/manifest.toml"))


def raw_dir(manifest: Manifest) -> Path:
    """Where the publication's raw files are kept: data/psgc/raw/<id>/, beside its folder."""
    return manifest.folder.parent / "raw" / manifest.id


def sha256_of(path: Path) -> str:
    """The SHA-256 of the file at path, as hex."""
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def cell_text(value) -> str:
    """A cell as the CSV writes it: text as it is, a whole number without .0, nothing for an empty cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _sheet_rows(workbook, sheet: str) -> list[list[str]]:
    """Every row of the sheet as text, padded to its widest row, with trailing empty rows dropped."""
    rows = [[cell_text(v) for v in row] for row in workbook[sheet].iter_rows(values_only=True)]
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    rows = [(r + [""] * width)[:width] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _write_csv(path: Path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerows(rows)


def _is_section(row) -> bool:
    """A Summary of Changes section header: a dated title alone in the first cell."""
    return bool(row[0]) and not any(row[1:]) and re.search(r"\d{4}", row[0]) is not None


def soc_rows(workbook, sheet: str, sections) -> list[list[str]]:
    """The rows of the named sections of a Summary of Changes sheet, each led by its section's name.

    Column headers and empty rows are skipped. Raises PsgcDataError if a section isn't in the sheet,
    or if a row of one has a value past the seven columns.
    """
    wanted, found, out, current = set(sections), set(), [], None
    for row in workbook[sheet].iter_rows(values_only=True):
        text = [cell_text(v) for v in row]
        head, rest = (text + [""] * 7)[:7], text[7:]
        if _is_section(head) and not any(rest):
            current = head[0] if head[0] in wanted else None
            found.update({current} - {None})
            continue
        if current is None or not any(text) or head[0] == SOC_COLUMNS[0]:
            continue
        if any(rest):
            raise PsgcDataError(f"{sheet}: section {current!r} has a value past column 7: {text}")
        out.append([current] + head)
    missing = [s for s in sections if s not in found]
    if missing:
        raise PsgcDataError(f"{sheet}: section {missing[0]!r} not found")
    return out


def extract(manifest: Manifest, raw=None, out=None):
    """Write the publication's CSVs into out (its folder by default) from the raw files in raw.

    Checks each raw file's SHA-256 against the manifest first, raising PsgcDataError naming the file
    if it differs.
    """
    raw = raw_dir(manifest) if raw is None else Path(raw)
    out = manifest.folder if out is None else Path(out)
    for f in manifest.files:
        actual = sha256_of(raw / f.name)
        if actual != f.sha256:
            raise PsgcDataError(f"{raw / f.name}: SHA-256 is {actual}, the manifest says {f.sha256}")
    for f in manifest.files:
        workbook = openpyxl.load_workbook(raw / f.name, read_only=True, data_only=True)
        try:
            if f.role == "datafile":
                for key, csv_name in SHEETS["datafile"].items():
                    _write_csv(out / csv_name, _sheet_rows(workbook, f.sheets[key]))
            else:
                rows = soc_rows(workbook, f.sheets["log"], manifest.soc_sections)
                _write_csv(out / SHEETS[f.role]["log"], [["section", *SOC_COLUMNS]] + rows)
        finally:
            workbook.close()


def check(root=None) -> tuple[list[str], list[str]]:
    """Re-extract every publication whose raw files are present and compare with the committed files.

    Returns the problems (a differing file, a hash mismatch) and the ids skipped for want of raw files.
    """
    problems, skipped = [], []
    for pub_id in publications(root):
        manifest = load_manifest(pub_id, root)
        if not all((raw_dir(manifest) / f.name).exists() for f in manifest.files):
            skipped.append(pub_id)
            continue
        with tempfile.TemporaryDirectory() as tmp:
            try:
                extract(manifest, out=tmp)
            except PsgcDataError as e:
                problems.append(f"{pub_id}: {e}")
                continue
            for f in manifest.files:
                for csv_name in SHEETS[f.role].values():
                    committed = manifest.folder / csv_name
                    if not committed.exists() or not filecmp.cmp(Path(tmp) / csv_name, committed, shallow=False):
                        problems.append(f"{pub_id}: {csv_name} differs from its raw file {f.name}")
    return problems, skipped
